#!/usr/bin/env python3
"""Quality-control pass over the completed cue sheet (revision 2).

Compares the completed workbook against the untouched template part by
part and cell by cell.  Revision 2 intentionally changes exactly three
package parts (sheet1 values/formula-result caches, appended shared
strings, and the fullCalcOnLoad flag); everything else must be identical.

In addition to the structural comparison, this checks the actual cue-sheet
content, the shared-string accounting, the calculation-chain references,
and -- critically for the earlier "doesn't work" report -- the cached
formula results that data-only viewers display before recalculation.
"""
import hashlib
import sys
import warnings
import zipfile
from collections import OrderedDict
from xml.etree import ElementTree as ET

import openpyxl

warnings.filterwarnings("ignore")

SRC = "Cue_Sheet_Template_2016V3-6-5.xlsx"
DST = "Cue_Sheet_The_Dark_Awaits_COMPLETED.xlsx"
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

fails = []


def check(label, ok, detail=""):
    print(("PASS  " if ok else "FAIL  ") + label + ((" :: " + detail) if detail else ""))
    if not ok:
        fails.append(label)


def sha256(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def sheet_cells(xml_bytes):
    """Ordered {coordinate: cell element} for a worksheet part."""
    root = ET.fromstring(xml_bytes)
    cells = OrderedDict()
    for cell in root.iter(NS + "c"):
        coord = cell.attrib.get("r")
        if coord:
            if coord in cells:
                raise AssertionError("duplicate cell %s" % coord)
            cells[coord] = cell
    return cells


def strings_from_sst(xml_bytes):
    root = ET.fromstring(xml_bytes)
    out = []
    for si in root.findall(NS + "si"):
        out.append("".join(t.text or "" for t in si.iter(NS + "t")))
    return root.attrib, out


# ---------------------------------------------------------------- 1. package
za, zb = zipfile.ZipFile(SRC), zipfile.ZipFile(DST)
check("template SHA-256 recorded", len(sha256(SRC)) == 64, sha256(SRC))
check("completed ZIP integrity (testzip)", zb.testzip() is None)
check("no duplicate ZIP part names",
      len(zb.namelist()) == len(set(zb.namelist())), "%d parts" % len(zb.namelist()))
check("package part list and order identical",
      za.namelist() == zb.namelist(), "%d parts" % len(zb.namelist()))
check("[Content_Types].xml remains first part",
      zb.namelist()[0] == "[Content_Types].xml")

changed = sorted(n for n in za.namelist() if za.read(n) != zb.read(n))
expected_changed = ["xl/sharedStrings.xml", "xl/workbook.xml",
                    "xl/worksheets/sheet1.xml"]
check("only sheet1.xml + sharedStrings.xml + workbook.xml changed",
      changed == expected_changed, "changed=%s" % changed)

for n in zb.namelist():
    if n.endswith(".xml") or n.endswith(".rels"):
        try:
            ET.fromstring(zb.read(n))
        except ET.ParseError as e:
            check("XML well-formed: %s" % n, False, str(e))
            break
else:
    check("all XML/rels parts well-formed", True)

wa = za.read("xl/workbook.xml").decode()
wb_xml = zb.read("xl/workbook.xml").decode()
check("workbook.xml diff is only fullCalcOnLoad",
      wb_xml.replace('<calcPr calcId="152511" fullCalcOnLoad="1"/>',
                     '<calcPr calcId="152511"/>') == wa)
check("workbook requests full calculation on load",
      'fullCalcOnLoad="1"' in wb_xml)

for part in ("[Content_Types].xml", "_rels/.rels", "xl/_rels/workbook.xml.rels",
             "xl/styles.xml", "xl/worksheets/sheet2.xml", "xl/calcChain.xml",
             "xl/theme/theme1.xml", "xl/tables/table1.xml",
             "xl/tables/table2.xml", "xl/tables/table3.xml",
             "xl/tables/table4.xml", "xl/tables/table5.xml",
             "xl/tables/table6.xml", "xl/tables/table7.xml",
             "xl/tables/table8.xml", "xl/printerSettings/printerSettings1.bin",
             "xl/printerSettings/printerSettings2.bin"):
    check("untouched: %s" % part, za.read(part) == zb.read(part))

# ------------------------------------------------------- 2. shared strings
attrs_a, strings_a = strings_from_sst(za.read("xl/sharedStrings.xml"))
attrs_b, strings_b = strings_from_sst(zb.read("xl/sharedStrings.xml"))
expected_new_strings = [
    "The Dark Awaits", "Dark Echo Productions",
    "Impact Wave Tick Interlude", "ZAZIE", "PRODUCTIONS",
    "Zazie Productions Publishing", "01372389230",
    "John's Story \u2014 Distorted Atmospheric Pad",
    "Spirits Still Around", "Industrial Electronic Interlude",
    "Mystic Bell Ambience", "Family Investigation \u2014 Dark Ambience",
]
check("template shared strings all preserved in original order",
      strings_b[:len(strings_a)] == strings_a,
      "%d original strings" % len(strings_a))
check("only the expected 12 unique strings appended",
      strings_b[len(strings_a):] == expected_new_strings,
      str(strings_b[len(strings_a):]))
check("sharedStrings uniqueCount equals <si> count",
      int(attrs_b["uniqueCount"]) == len(strings_b) == 302,
      "uniqueCount=%s si=%d" % (attrs_b["uniqueCount"], len(strings_b)))

cells_a1 = sheet_cells(za.read("xl/worksheets/sheet1.xml"))
cells_b1 = sheet_cells(zb.read("xl/worksheets/sheet1.xml"))
cells_b2 = sheet_cells(zb.read("xl/worksheets/sheet2.xml"))
shared_refs = []
inline_cells = []
for cells in (cells_b1, cells_b2):
    for coord, cell in cells.items():
        if cell.attrib.get("t") == "s":
            v = cell.find(NS + "v")
            shared_refs.append(int(v.text))
        if cell.attrib.get("t") == "inlineStr":
            inline_cells.append(coord)
check("sharedStrings count equals actual t=\"s\" cells in both sheets",
      int(attrs_b["count"]) == len(shared_refs) == 378,
      "count=%s actual=%d" % (attrs_b["count"], len(shared_refs)))
check("all shared-string indexes in range",
      all(0 <= i < len(strings_b) for i in shared_refs),
      "min=%d max=%d" % (min(shared_refs), max(shared_refs)))
check("no inlineStr cells remain (Excel-native shared strings used)",
      not inline_cells and b"inlineStr" not in zb.read("xl/worksheets/sheet1.xml"),
      str(inline_cells[:5]))

# ------------------------------------------- 3. raw worksheet preservation
check("sheet1 cell coordinates and XML order identical",
      list(cells_a1.keys()) == list(cells_b1.keys()),
      "%d cells" % len(cells_b1))
check("every sheet1 cell keeps its exact style index s=",
      all(cells_a1[c].attrib.get("s") == cells_b1[c].attrib.get("s")
          for c in cells_a1))
check("no sheet1 cell gained/lost any attribute except formula-cache/input type t=",
      all(set(cells_a1[c].attrib).symmetric_difference(set(cells_b1[c].attrib))
          <= {"t"} for c in cells_a1))


def formula_map(cells):
    out = {}
    for coord, cell in cells.items():
        f = cell.find(NS + "f")
        if f is not None:
            out[coord] = (f.text or "", tuple(sorted(f.attrib.items())))
    return out


fa_raw, fb_raw = formula_map(cells_a1), formula_map(cells_b1)
check("raw <f> formula text/attributes identical for every formula cell",
      fa_raw == fb_raw, "template=%d completed=%d" % (len(fa_raw), len(fb_raw)))

cc_root = ET.fromstring(zb.read("xl/calcChain.xml"))
cc_refs = [c.attrib["r"] for c in cc_root.findall(NS + "c")]
check("calcChain references exactly the unchanged formula-cell set",
      set(cc_refs) == set(fb_raw) and len(cc_refs) == len(fb_raw),
      "calcChain=%d formulas=%d" % (len(cc_refs), len(fb_raw)))
check("calcChain still points at sheetId 13 (Template)",
      all(c.attrib.get("i") == "13" for c in cc_root.findall(NS + "c")))

# ---------------------------------------------------------------- 4. structure
a = openpyxl.load_workbook(SRC)
b = openpyxl.load_workbook(DST)
b_cached = openpyxl.load_workbook(DST, data_only=True)

check("sheet names/state identical",
      [(s.title, s.sheet_state) for s in a.worksheets] ==
      [(s.title, s.sheet_state) for s in b.worksheets],
      str([(s.title, s.sheet_state) for s in b.worksheets]))
check("hidden 'hidat' worksheet still hidden", b["hidat"].sheet_state == "hidden")

ta, tb, tc = a["Template"], b["Template"], b_cached["Template"]
check("dimensions identical", ta.dimensions == tb.dimensions == "A1:R999")
check("data validations identical",
      [str(d.sqref) + "|" + str(d.formula1) + "|" + d.type
       for d in ta.data_validations.dataValidation] ==
      [str(d.sqref) + "|" + str(d.formula1) + "|" + d.type
       for d in tb.data_validations.dataValidation],
      "%d dropdown/range rules kept" % len(tb.data_validations.dataValidation))
check("conditional formatting identical",
      len(ta.conditional_formatting._cf_rules) == len(tb.conditional_formatting._cf_rules))
check("merged cells identical",
      sorted(str(r) for r in ta.merged_cells.ranges) ==
      sorted(str(r) for r in tb.merged_cells.ranges))
check("column widths/hidden identical",
      {k: (v.width, v.hidden) for k, v in ta.column_dimensions.items()} ==
      {k: (v.width, v.hidden) for k, v in tb.column_dimensions.items()})
check("sheet protection preserved",
      (ta.protection.sheet, ta.protection.algorithmName) ==
      (tb.protection.sheet, tb.protection.algorithmName))
check("workbook structure protection preserved",
      a.security.lockStructure == b.security.lockStructure,
      "lockStructure=%s" % b.security.lockStructure)
check("CueDetails2 table range preserved",
      str(tb.tables["CueDetails2"].ref) == "A20:R999")
check("CueDetails2 calculated-column formulas preserved",
      [c.calculatedColumnFormula.attr_text if c.calculatedColumnFormula else None
       for c in ta.tables["CueDetails2"].tableColumns] ==
      [c.calculatedColumnFormula.attr_text if c.calculatedColumnFormula else None
       for c in tb.tables["CueDetails2"].tableColumns])


def openpyxl_formulas(ws):
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                out[c.coordinate] = c.value
    return out


fa, fb = openpyxl_formulas(ta), openpyxl_formulas(tb)
check("openpyxl formula view identical (%d formulas)" % len(fa), fa == fb)

# ---------------------------------------------------------------- 5. content
def val(ws, coord):
    return ws[coord].value


expect_header = {
    "N4": "The Dark Awaits",
    "N11": "Dark Echo Productions",
    "D13": 45,
    "F13": 0,
    "C4": "Cue Sheet Classification:",
    "M4": "Program (series, film, etc.) Title:",
    "M8": "Episode Number:",
    "M11": "Production Company:",
}
for coord, want in expect_header.items():
    got = val(tb, coord)
    check("%s = %r" % (coord, want), got == want, "got %r" % got)

for coord in ("D4", "D5", "D8", "D9", "D10", "D11",
              "N5", "N6", "N7", "N8", "N9", "N12", "N13", "N14"):
    got = val(tb, coord)
    check("%s left blank (unknown / instructed blank)" % coord,
          got in (None, ""), "got %r" % got)

expect_cached_header = {
    "A1": "The Dark Awaits", "A2": "Music Cue Sheet",
    "D14": 25, "E14": "min.", "F14": 8, "G14": "sec.", "H14": "(auto calc.)",
}
for coord, want in expect_cached_header.items():
    got = val(tc, coord)
    check("cached formula result %s = %r" % (coord, want), got == want,
          "got %r" % got)

CUES = [
    ("Impact Wave Tick Interlude", 0, 2, 32, 0, 2, 35),
    ("John's Story \u2014 Distorted Atmospheric Pad", 0, 2, 42, 0, 2, 48),
    ("Spirits Still Around", 0, 4, 6, 0, 4, 25),
    ("Industrial Electronic Interlude", 0, 6, 12, 0, 6, 49),
    ("Mystic Bell Ambience", 0, 7, 32, 0, 8, 0),
    ("Family Investigation \u2014 Dark Ambience", 0, 14, 25, 0, 38, 0),
]

print("\n--- cue blocks as written (grouped template layout) ---")
for i, (title, ih, im, isec, oh, om, osec) in enumerate(CUES):
    r, pr = 20 + i * 2, 21 + i * 2
    dur = (oh * 3600 + om * 60 + osec) - (ih * 3600 + im * 60 + isec)
    first_ok = (
        val(tb, "B%d" % r) == title and val(tb, "C%d" % r) == "BI"
        and tuple(val(tb, "%s%d" % (c, r)) for c in "DEF") == (ih, im, isec)
        and tuple(val(tb, "%s%d" % (c, r)) for c in "GHI") == (oh, om, osec)
        and val(tb, "L%d" % r) == "Composer"
        and val(tb, "M%d" % r) == "ZAZIE"
        and val(tb, "N%d" % r) == "PRODUCTIONS"
        and val(tb, "O%d" % r) is None and val(tb, "P%d" % r) is None
        and val(tb, "Q%d" % r) == "BMI" and val(tb, "R%d" % r) == 1
        and val(tc, "A%d" % r) == i + 1
        and val(tc, "J%d" % r) == dur // 60
        and val(tc, "K%d" % r) == dur % 60)
    check("cue %d first row %d (Composer): %s" % (i + 1, r, title), first_ok,
          "seq(cache)=%r dur(cache)=%r:%r writer=%r %r"
          % (val(tc, "A%d" % r), val(tc, "J%d" % r), val(tc, "K%d" % r),
             val(tb, "M%d" % r), val(tb, "N%d" % r)))
    pub_ok = (
        all(val(tb, "%s%d" % (c, pr)) is None for c in "BCDEFGHI")
        and val(tc, "A%d" % pr) is None
        and val(tc, "J%d" % pr) is None and val(tc, "K%d" % pr) is None
        and val(tb, "L%d" % pr) == "Publisher"
        and val(tb, "M%d" % pr) is None and val(tb, "N%d" % pr) is None
        and val(tb, "O%d" % pr) == "Zazie Productions Publishing"
        and val(tb, "P%d" % pr) == "01372389230"
        and isinstance(val(tb, "P%d" % pr), str)
        and val(tb, "Q%d" % pr) == "BMI" and val(tb, "R%d" % pr) == 1)
    check("cue %d publisher row %d carries party only (B:K blank)" % (i + 1, pr),
          pub_ok,
          "O=%r P=%r Q=%r R=%r" % (val(tb, "O%d" % pr), val(tb, "P%d" % pr),
                                   val(tb, "Q%d" % pr), val(tb, "R%d" % pr)))
    print("      seq %d rows %d-%d | %-42s | BI | %d:%02d:%02d -> %d:%02d:%02d "
          "| %d:%02d | PRODUCTIONS, ZAZIE (BMI) / Zazie Productions Publishing (BMI)"
          % (i + 1, r, pr, title, ih, im, isec, oh, om, osec, dur // 60, dur % 60))

times = [(ih * 3600 + im * 60 + isec, oh * 3600 + om * 60 + osec)
         for (_, ih, im, isec, oh, om, osec) in CUES]
total = sum(o - i for i, o in times)
check("every cue has a positive duration", all(o > i for i, o in times),
      "; ".join("%ds" % (o - i) for i, o in times))
check("cues in chronological order (non-decreasing time-in)",
      all(times[k][0] >= times[k - 1][0] for k in range(1, len(times))))
check("total music duration is 25 min. 8 sec. and fits the 45-minute program",
      total == 1508 and val(tc, "D14") * 60 + val(tc, "F14") == 1508
      and total <= 45 * 60, "%ds" % total)
check("all six cues use Usage Code BI",
      all(val(tb, "C%d" % (20 + i * 2)) == "BI" for i in range(6)))
check("writer on every cue is PRODUCTIONS, ZAZIE / BMI",
      all(val(tb, "N%d" % r) + ", " + val(tb, "M%d" % r) == "PRODUCTIONS, ZAZIE"
          and val(tb, "Q%d" % r) == "BMI" for r in range(20, 32, 2)))
check("publisher on every cue is Zazie Productions Publishing / BMI / IPI 01372389230",
      all(val(tb, "O%d" % r) == "Zazie Productions Publishing"
          and val(tb, "Q%d" % r) == "BMI"
          and val(tb, "P%d" % r) == "01372389230" for r in range(21, 32, 2)))
check("shares display as 100.00% (stored 1, format 0.00%) on all party rows",
      all(val(tb, "R%d" % r) == 1 and tb["R%d" % r].number_format == "0.00%"
          for r in range(20, 32)))
check("BMI account 4380631 was not entered into any IPI cell",
      all(str(val(tb, "P%d" % r)) != "4380631" for r in range(20, 32)))
check("writer IPI 01280391067 not forced into the Publisher IPI column",
      all(val(tb, "P%d" % r) != "01280391067" for r in range(20, 32)))

hidat = b["hidat"]
usage_list = [hidat["E%d" % r].value for r in range(2, 9)]
role_list = [hidat["F%d" % r].value for r in range(2, 7)
             if hidat["F%d" % r].value]
affiliation_list = [hidat["H%d" % r].value for r in range(2, 178)]
check("entered dropdown values are exact members of hidat lists",
      "BI" in usage_list and "Composer" in role_list
      and "Publisher" in role_list and "BMI" in affiliation_list,
      "usage=%s roles=%s" % (usage_list, role_list))
check("entered times satisfy the template's whole-number validations",
      all(0 <= hh <= 9 and 0 <= mm <= 59 and 0 <= ss <= 59
          for (_, hh, mm, ss, hh2, mm2, ss2) in CUES
          for hh, mm, ss in ((hh, mm, ss), (hh2, mm2, ss2))))

dirty = []
for r in range(32, 200):
    for c in "BCDEFGHILMNOPQR":
        if val(tb, "%s%d" % (c, r)) not in (None, ""):
            dirty.append("%s%d" % (c, r))
check("no stray data below the cue block", not dirty, str(dirty[:8]))
check("filled input cells keep the template's unlocked style",
      all(tb["%s%d" % (c, r)].protection.locked is False
          for r in range(20, 32) for c in "BCDEFGHILMNOPQR"
          if val(tb, "%s%d" % (c, r)) not in (None, "")))
check("E/F/H/I time cells keep 00 format on cue first rows",
      all(tb["%s%d" % (c, r)].number_format == "00"
          for r in range(20, 32, 2) for c in "EFHI"))

print()
print("RESULT:", "ALL CHECKS PASSED" if not fails
      else "%d FAILURES: %s" % (len(fails), fails))
sys.exit(1 if fails else 0)
